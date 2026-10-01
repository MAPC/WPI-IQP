"""Verify delivered artifacts and invariants without relying on file existence alone."""
from pathlib import Path
import csv
import json
import zipfile
import pandas as pd
import pyogrio
from openpyxl import load_workbook
from .utils import sha256, write_json, html_page

def run_acceptance(root,assets,sites,layers,tests,baseline=None):
    root=Path(root);checks=[]
    def add(name,passed,evidence):checks.append({"requirement":name,"passed":bool(passed),"evidence":evidence})
    with (root/"bootstrap/root_inventory.csv").open(encoding="utf-8-sig") as f:
        raw=list(csv.DictReader(f))
    present=[r for r in raw if (root/r["filename"]).is_file()]
    intact=all(sha256(root/r["filename"])==r["sha256"] for r in present)
    add("Available original root files preserved",intact,f"SHA-256 verified for {len(present)} original files. Duplicate root copies are optional on a new handoff computer.")
    parsed_source=assets.source_file.dropna().iloc[0]
    parsed_hash=assets.source_sha256.dropna().iloc[0]
    add("Authoritative input unchanged during processing",sha256(root/parsed_source)==parsed_hash,"Current input hash equals the source checksum retained on every parsed row; future replacement exports need not equal bootstrap's initial version.")
    required=["data/assets_clean.csv","data/assets_clean.xlsx","data/sites_summary.csv","data/sites_summary.xlsx","data/map_dataset.csv","data/attribute_dictionary.csv","data/amenity_dictionary.csv","data/activity_dictionary.csv","data/factcheck_audit.csv","data/transit_discrepancies.csv","analysis/analysis_workbook.xlsx","documentation/data_dictionary.xlsx","documentation/methodology.md","gis/mapc_networks.gpkg","gis/bicycle_facilities_web.geojson","gis/shared_use_paths_web.geojson","gis/walking_trails_web.geojson","maps/MAPC_access_map.html","maps/MAPC_access_map_static.png","maps/MAPC_access_map_static.svg","reports/schema_report.html"]
    for relative in required:
        p=root/"output"/relative;add("Artifact: "+relative,p.is_file() and p.stat().st_size>0,"File exists and is nonempty")
    clean=pd.read_csv(root/"output/data/assets_clean.csv")
    add("Clean asset export row reconciliation",len(clean)==len(assets),f"{len(clean)} saved vs {len(assets)} parsed source records")
    add("Asset source traceability",clean.source_row_uid.is_unique and clean.source_row_number.notna().all(),"Unique snapshot row UIDs and CSV record locators retained")
    summary=pd.read_csv(root/"output/data/sites_summary.csv")
    add("Site export reconciliation",len(summary)==len(sites) and summary.asset_count.sum()==assets.record_status.eq("accepted").sum(),f"{len(summary)} sites; summed members equal accepted assets")
    for p in sorted((root/"output").rglob("*.xlsx")):
        with zipfile.ZipFile(p) as z:healthy=z.testzip() is None
        wb=load_workbook(p,read_only=True,data_only=False)
        add("Workbook opens: "+p.name,healthy and bool(wb.sheetnames),f"{len(wb.sheetnames)} readable worksheets; OOXML ZIP intact")
        wb.close()
    gpkg=root/"output/gis/mapc_networks.gpkg"
    saved_layers={row[0] for row in pyogrio.list_layers(gpkg)}
    for name,frame in layers.items():
        info=pyogrio.read_info(gpkg,layer=name)
        add("Full GIS layer: "+name,name in saved_layers and info["features"]==len(frame),f"{info['features']} features; CRS {info['crs']}; source full geometry count {len(frame)}")
        if name!="land_line_systems":
            content=json.loads((root/f"output/gis/{name}_web.geojson").read_text(encoding="utf-8"))
            eligible=int(frame.analysis_include.eq(True).sum())
            add("GIS current/public web filter: "+name,len(content["features"])==eligible,f"{len(content['features'])} web features vs {eligible} eligible full-resolution features")
    charts=list((root/"output/charts").glob("*.png"))
    add("Publication figures",len(charts)>=6 and all(p.with_suffix(".svg").exists() for p in charts),f"{len(charts)} PNG/SVG figure pairs")
    add("Unit and integration tests",tests["unit_and_integration"]["status"]=="passed",tests["unit_and_integration"].get("summary",tests["unit_and_integration"]["status"]))
    browser=tests["map"]["status"]
    add("Browser smoke test or documented environment limitation",browser in ("passed","unavailable"),browser)
    if browser=="passed":add("Browser screenshots",(root/"output/reports/map_smoke_test.png").exists() and (root/"output/reports/map_offline_test.png").exists(),"Desktop and offline screenshots saved; online tile verification is reported separately by the browser check")
    if baseline:
        rebuild=json.loads((root/"output/reports/clean_rebuild_acceptance.json").read_text())
        add("Clean rebuild canonical CSV reproduction",rebuild["passed"],f"{len(rebuild['checks'])} canonical files compared against preserved baseline")
    result={"passed":all(c["passed"] for c in checks),"clean_rebuild_performed":bool(baseline),"checks":checks,"remaining_delivery_step":"Create release ZIP after successful manifest and QC generation."}
    write_json(root/"output/reports/acceptance_report.json",result)
    (root/"output/reports/acceptance_report.html").write_text(html_page("MAPC acceptance checks",pd.DataFrame(checks).to_html(index=False,escape=True)),encoding="utf-8")
    return result
