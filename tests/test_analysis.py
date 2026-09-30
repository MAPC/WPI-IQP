"""Tests for scientific denominators, clustered units and typed exports."""
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.analysis import (validated_population, cross_states, site_sensitivity,
                          correlation_table, frequency_table, run_analysis)
from src.documentation import build_data_dictionary
from src.workbook import Book, cell_value


def example_assets():
    rows = []
    for i, (site, validated, transit, parking, attrs, acts) in enumerate([
        ("s1", True, "YES", "UNKNOWN", ["Accessible Parking", "Family Friendly"], ["Walking"]),
        ("s1", True, "NO", "YES", ["Family Friendly", "Restrooms Available"], ["Walking", "Picnicking"]),
        ("s2", True, "UNKNOWN", "NO", [], []),
        ("s1", False, "YES", "YES", ["Unfinished feature"], ["Unfinished activity"]),
        ("s3", None, "UNKNOWN", "UNKNOWN", [], []),
    ]):
        rows.append({"asset_id": f"a{i}", "source_record_key": f"a{i}", "asset_name": f"Asset {i}",
                     "site_id": site, "site_name": site, "municipality": "Town", "subregion": "Inner Core",
                     "field_validated": validated, "staff_reviewed": False, "coordinate_usable": True,
                     "record_status": "accepted", "attribute_list": attrs, "attribute_count": len(attrs),
                     "amenity_list": [], "amenity_count": 0, "activity_list": acts, "activity_count": len(acts),
                     "near_public_transit": transit, "free_entry_parking": parking,
                     "accessible_parking": "YES" if i == 0 else "UNKNOWN", "accessible_restroom": "UNKNOWN",
                     "wheelchair_stroller_friendly_trail": "UNKNOWN", "restrooms_available": "UNKNOWN",
                     "transportation_access_profile": "Free parking only" if i == 1 else "Unknown / unresolved"})
    return pd.DataFrame(rows)


def test_validated_universe_rejects_quarantine_and_unfinished_and_deduplicates():
    frame = example_assets()
    frame = pd.concat([frame, frame.iloc[[0]]], ignore_index=True)
    frame.loc[2, "record_status"] = "quarantined"
    result = validated_population(frame)
    assert set(result.asset_id) == {"a0", "a1"}


def test_unknown_is_not_no_and_all_denominator_reconciles():
    frame = validated_population(example_assets())
    result = cross_states(frame, "near_public_transit", "free_entry_parking")
    assert result.n.sum() == len(frame)
    assert result.denominator_all_validated.eq(3).all()
    assert result.loc[result.state_a.eq("YES") & result.state_b.eq("UNKNOWN"), "n"].iloc[0] == 1
    assert result.loc[result.state_a.eq("YES") & result.state_b.eq("NO"), "n"].iloc[0] == 0
    assert result.share_all_validated.sum() == pytest.approx(1)


def test_site_union_uses_validated_members_and_not_asset_sum():
    sites = site_sensitivity(validated_population(example_assets())).set_index("site_id")
    assert sites.at["s1", "n_validated_assets"] == 2
    assert sites.at["s1", "attribute_count"] == 3
    assert sites.at["s1", "activity_count"] == 2
    assert sites.at["s1", "mean_asset_activity_count"] == 1.5
    assert sites.at["s1", "transportation_access_profile"] == "Transit + free parking"
    assert sites.at["s2", "transportation_access_profile"] == "Unknown / unresolved"


def test_spearman_ties_and_constant_input_are_handled():
    frame = pd.DataFrame({"attribute_count": [1, 1, 3], "activity_count": [2, 2, 4], "amenity_count": [0, 0, 0]})
    corr = correlation_table(frame)
    assert corr.iloc[0].spearman_rho == pytest.approx(1)
    assert pd.isna(corr.iloc[1].spearman_rho)
    assert corr.n_complete_pairs.eq(3).all()


def test_activity_frequency_counts_each_asset_once():
    frame = pd.DataFrame({"activity_list": [["Walking", "Walking"], ["Walking", "Fishing"], []]})
    frequency = frequency_table(frame, "activity_list", "activity").set_index("activity")
    assert frequency.at["Walking", "n_assets"] == 2
    assert frequency.at["Walking", "share_validated_assets"] == pytest.approx(2 / 3)


def test_analysis_generates_all_questions_and_explicit_access_denominators(tmp_path):
    tables = run_analysis(example_assets(), pd.DataFrame(), tmp_path, {})
    assert tables["findings"].question_number.tolist() == list(range(1, 14))
    assert tables["quality"].set_index("metric").at["Validated unique assets analyzed", "count"] == 3
    access = tables["accessibility"].set_index("field")
    assert access.at["accessible_parking", "n_known"] == 1
    assert access.at["accessible_parking", "yes_share_known"] == 1
    assert access.at["accessible_parking", "yes_share_all"] == pytest.approx(1 / 3)
    assert (tmp_path / "output/analysis/findings.md").exists()
    assert len(tables["validated_sites"]) == 2
    assert pd.api.types.is_numeric_dtype(tables["accessibility_richness"].accessibility_yes_count)


def test_empty_validated_analysis_does_not_invent_zero_rates(tmp_path):
    frame = example_assets()
    frame["field_validated"] = False
    tables = run_analysis(frame, pd.DataFrame(), tmp_path, {})
    assert tables["accessibility"].yes_share_all.isna().all()
    assert len(tables["findings"]) == 13


def test_excel_preserves_numeric_values_leading_zero_id_and_safe_text(tmp_path):
    path = tmp_path / "typed.xlsx"
    book = Book(path)
    sheet = book.sheet("Values")
    book.table(sheet, pd.DataFrame({"zip_code": ["02110"], "count": [3], "share": [.25],
                                   "source_text": ["=2+2"], "activity_list": [["Walking", "Fishing"]]}))
    assert book.close()["status"] == "passed"
    saved = load_workbook(path, data_only=False)
    row = list(saved["Values"].iter_rows(min_row=2, max_row=2))[0]
    assert row[0].value == "02110"
    assert row[1].value == 3 and row[1].data_type == "n"
    assert row[2].value == .25 and row[2].data_type == "n"
    assert row[3].value == "=2+2" and row[3].data_type == "s"
    assert json.loads(row[4].value) == ["Walking", "Fishing"]
    saved.close()


def test_nullable_numeric_distances_export_as_blank_not_invalid_number(tmp_path):
    path = tmp_path / "nullable.xlsx"
    book = Book(path)
    sheet = book.sheet("Distances")
    book.table(sheet, pd.DataFrame({"nearest_transit_distance_miles": [1.25, float("nan"), None],
                                   "attribute_count": [2, 0, 1], "nearest_shared_use_path_name": ["Path A", None, "Path C"],
                                   "near_shared_use_path": ["YES", "UNKNOWN", "NO"]}), data_sheet=True)
    assert book.close()["status"] == "passed"
    saved = load_workbook(path, data_only=True)
    assert saved["Distances"]["A2"].value == 1.25
    assert saved["Distances"]["A3"].value is None
    assert saved["Distances"]["A4"].value is None
    saved.close()


def test_dictionary_covers_every_asset_and_site_field(tmp_path):
    frame = example_assets()
    sites = site_sensitivity(validated_population(frame))
    dictionary = build_data_dictionary(frame, sites, {}, tmp_path)
    assert set(dictionary[dictionary.dataset.eq("assets_clean")].field_name) == set(frame.columns)
    assert set(dictionary[dictionary.dataset.eq("sites_summary")].field_name) == set(sites.columns)
    assert dictionary.description.str.len().gt(15).all()
    assert dictionary.derivation_rule.str.len().gt(15).all()
